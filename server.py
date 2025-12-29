from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
import csv
import copy
import os
import tempfile
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from werkzeug.utils import secure_filename

app = Flask(__name__)
CORS(app)  # Enable CORS for the React frontend

def duplicate_slide(pres, index):
    """
    Duplicate the slide at the given index in the presentation.
    """
    source = pres.slides[index]
    blank_slide_layout = pres.slide_layouts[6]
    dest = pres.slides.add_slide(blank_slide_layout)

    for shape in source.shapes:
        newel = copy.deepcopy(shape.element)
        dest.shapes._spTree.insert_element_before(newel, 'p:extLst')

    return pres.slides[-1]

def replace_text_in_shape(shape, name, org_country_text, initial_name_size=60, initial_org_size=30):
    """
    Replace text in shape with proper formatting.
    """
    if not shape.has_text_frame:
        return

    text_frame = shape.text_frame
    text_frame.clear()

    # Add name paragraph
    p_name = text_frame.paragraphs[0]
    p_name.text = name
    p_name.alignment = PP_ALIGN.CENTER

    if p_name.runs:
        for run in p_name.runs:
            run.font.name = 'Cambria'
            run.font.size = Pt(initial_name_size)
            run.font.bold = True

    # Add organization-country paragraph if provided
    if org_country_text:
        p_org = text_frame.add_paragraph()
        p_org.text = org_country_text
        p_org.alignment = PP_ALIGN.CENTER

        if p_org.runs:
            for run in p_org.runs:
                run.font.name = 'Cambria'
                run.font.size = Pt(initial_org_size)
                run.font.bold = False

    # Enable auto-fit
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

@app.route('/generate', methods=['POST'])
def generate_nameplates():
    try:
        # Get uploaded files
        csv_file = request.files.get('csv')
        pptx_file = request.files.get('pptx')
        
        # Get options
        include_org = request.form.get('includeOrg') == 'true'
        include_country = request.form.get('includeCountry') == 'true'
        
        # Check for template.pptx in same folder
        use_local_template = False
        if os.path.exists('template.pptx'):
            use_local_template = True
            print("Using local template.pptx file")
        
        if not csv_file:
            return jsonify({'error': 'Missing CSV file'}), 400
        
        if not use_local_template and not pptx_file:
            return jsonify({'error': 'Missing PowerPoint template'}), 400

        # Create temporary directory for processing
        with tempfile.TemporaryDirectory() as temp_dir:
            # Save uploaded files
            csv_path = os.path.join(temp_dir, secure_filename(csv_file.filename))
            csv_file.save(csv_path)
            
            # Use local template or uploaded file
            if use_local_template:
                pptx_path = 'template.pptx'
            else:
                pptx_path = os.path.join(temp_dir, secure_filename(pptx_file.filename))
                pptx_file.save(pptx_path)
            
            output_path = os.path.join(temp_dir, 'Generated_Nameplates.pptx')

            # Parse CSV Data
            participants = []
            with open(csv_path, 'r', encoding='utf-8-sig') as f:
                first_line = f.readline()
                f.seek(0)

                # Detect delimiter
                comma_count = first_line.count(',')
                semicolon_count = first_line.count(';')
                tab_count = first_line.count('\t')

                if semicolon_count > comma_count and semicolon_count > tab_count:
                    delimiter = ';'
                elif tab_count > comma_count:
                    delimiter = '\t'
                else:
                    delimiter = ','

                reader = csv.DictReader(f, delimiter=delimiter)
                fieldnames = [name.strip() for name in reader.fieldnames if name]

                for row in reader:
                    normalized_row = {}
                    for k, v in row.items():
                        if k:
                            clean_key = k.strip()
                            clean_value = v.strip() if v else ''
                            normalized_row[clean_key] = clean_value

                    if any(normalized_row.values()):
                        participants.append(normalized_row)

            if not participants:
                return jsonify({'error': 'No participants found in CSV'}), 400

            # Load Presentation
            prs = Presentation(pptx_path)
            template_slide_index = 0

            # Generate slides
            for index, person in enumerate(participants):
                name = (person.get('Name') or person.get('name') or
                        person.get('NAME') or person.get('Nama') or '').strip()
                org = (person.get('Organization') or person.get('organisation') or
                       person.get('ORGANIZATION') or person.get('Org') or '').strip()
                country = (person.get('Country') or person.get('country') or
                           person.get('COUNTRY') or '').strip()

                if not name:
                    continue

                # Apply include options
                if not include_org:
                    org = ''
                if not include_country:
                    country = ''

                # Combine Org and Country
                if org and country:
                    org_country_text = f"{org} - {country}"
                elif org:
                    org_country_text = org
                elif country:
                    org_country_text = country
                else:
                    org_country_text = ""

                # Duplicate the template slide
                new_slide = duplicate_slide(prs, template_slide_index)

                # Collect all text shapes
                all_text_shapes = []
                for shape in new_slide.shapes:
                    if shape.has_text_frame:
                        shape.text_frame.clear()
                        all_text_shapes.append(shape)

                # Find main text shape (largest)
                main_text_shape = None
                max_area = 0
                for shape in all_text_shapes:
                    current_area = shape.width * shape.height
                    if current_area > max_area:
                        max_area = current_area
                        main_text_shape = shape

                if main_text_shape:
                    # Apply content to main shape
                    replace_text_in_shape(main_text_shape, name, org_country_text)

                    # Create top reversed shape
                    left = main_text_shape.left
                    width = main_text_shape.width
                    height = main_text_shape.height

                    slide_height = prs.slide_height
                    bottom_shape_bottom_edge = main_text_shape.top + main_text_shape.height
                    margin_from_bottom_edge = slide_height - bottom_shape_bottom_edge
                    new_top = margin_from_bottom_edge

                    top_reversed_shape = new_slide.shapes.add_textbox(left, new_top, width, height)
                    replace_text_in_shape(top_reversed_shape, name, org_country_text)
                    top_reversed_shape.rotation = 180

            # Remove template slide
            xml_slides = prs.slides._sldIdLst
            slides = list(xml_slides)
            xml_slides.remove(slides[0])

            # Save output
            prs.save(output_path)

            # Read file into memory before sending
            with open(output_path, 'rb') as f:
                file_data = f.read()

            # Create response with file data from memory
            from flask import Response
            response = Response(
                file_data,
                mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                headers={
                    'Content-Disposition': 'attachment; filename=Generated_Nameplates.pptx'
                }
            )
            
            return response

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/health', methods=['GET'])
def health_check():
    has_template = os.path.exists('template.pptx')
    return jsonify({
        'status': 'ok',
        'has_template': has_template
    })

if __name__ == '__main__':
    print("=" * 50)
    print("Nameplate Generator Server")
    print("=" * 50)
    print("\nServer is running on: http://localhost:5000")
    
    # Check for template file
    if os.path.exists('template.pptx'):
        print("\n✓ Found template.pptx - will use this automatically")
        print("  (You don't need to upload a template file)")
    else:
        print("\n⚠ No template.pptx found in folder")
        print("  (You'll need to upload a template each time)")
        print("  Tip: Place your template as 'template.pptx' in this folder")
    
    print("\nKeep this window open while using the app.")
    print("Press Ctrl+C to stop the server.")
    print("=" * 50)
    app.run(debug=True, port=5000)